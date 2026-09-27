#!/usr/bin/env python3
"""Generate Streamix / DDB project Word report."""
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "Streamix_Distributed_Movie_Streaming_Complete_Report.docx",
)

doc = Document()

# Narrow margins slightly for denser academic look
for section in doc.sections:
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def set_heading_color(paragraph, hex_color="B20710"):
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(hex_color)


def add_h(text, level=1):
    p = doc.add_heading(text, level=level)
    set_heading_color(p)
    return p


def add_p(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)
    return p


def add_bullet(text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level:
        p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    return p


def add_code(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.15)
    return p


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(10)
    for ri, row in enumerate(rows):
        cells = table.rows[ri + 1].cells
        for ci, val in enumerate(row):
            cells[ci].text = str(val)
            for p in cells[ci].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    doc.add_paragraph()
    return table


# ========== COVER ==========
for _ in range(2):
    doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("STREAMIX")
r.bold = True
r.font.size = Pt(36)
r.font.color.rgb = RGBColor(0xE5, 0x09, 0x14)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Distributed Movie Streaming Database System")
r.bold = True
r.font.size = Pt(18)

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run("Complete Project Report — Implementation, Architecture & DDB Concepts")
r.font.size = Pt(13)
r.italic = True

doc.add_paragraph()
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run(
    "Course Project: Distributed Databases\n"
    "Polyglot Microservices Backend + React Native (Netflix-style) Client\n"
    "Local free stack: PostgreSQL · MongoDB · Redis · Node.js / Express · Docker"
).font.size = Pt(11)

doc.add_paragraph()
note = doc.add_paragraph()
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
nr = note.add_run(
    "This document describes the system as implemented: what was built, "
    "how it maps to the architecture proposal, how to run and demo it, "
    "and how to explain Distributed Database concepts in a viva."
)
nr.font.size = Pt(10)
nr.italic = True

doc.add_page_break()

# ========== TOC-ish ==========
add_h("Table of Contents", 1)
toc_items = [
    "1. Executive Summary",
    "2. Project Objectives & Scope",
    "3. System Architecture Overview",
    "4. Technology Stack",
    "5. Distributed Database Concepts Implemented",
    "6. Polyglot Persistence Design",
    "7. Microservices & API Gateway",
    "8. Data Model & Fragmentation",
    "9. Shard Routing, Replication & Failover",
    "10. Region-Aware Routing & Near You",
    "11. Recommendations Engine",
    "12. Caching & Event Bus (Redis)",
    "13. Analytics Service",
    "14. Mobile Application (Streamix)",
    "15. Admin & Catalog Enrichment (TMDb)",
    "16. Repository Structure",
    "17. How to Run (Setup Guide)",
    "18. Demonstration Scripts (Viva)",
    "19. Alignment with Architecture Proposal",
    "20. Phase-2 / Future Work",
    "21. Conclusion & Viva Talking Points",
    "Appendix A — API Endpoints",
    "Appendix B — Docker Containers",
    "Appendix C — Key Source Files",
]
for t in toc_items:
    add_bullet(t)

doc.add_page_break()

# ========== 1 ==========
add_h("1. Executive Summary", 1)
add_p(
    "Streamix is a Netflix-style movie streaming catalog platform built as a "
    "course project in Distributed Databases. The system demonstrates a working "
    "backend that is not a single database — it is a distributed, polyglot "
    "cluster of services and stores — plus a React Native mobile client."
)
add_p(
    "The backend is implemented as microservices behind an API gateway on port 4000. "
    "PostgreSQL stores users (with primary–replica replication) and the movie catalog "
    "(horizontally sharded into US and EU regions, each with its own replica). "
    "MongoDB stores watchlists and reviews with application-level hashing across "
    "two shards each. Redis provides catalog response caching and a Streams-based "
    "event bus (stand-in for Kafka). Separate recommendation and analytics services "
    "complete the service layer."
)
add_p(
    "The mobile app (StreamixApp) provides Home (Trending, Recommended, Near You, "
    "genre rows), Search, Watchlist, Movie Detail with in-app YouTube trailers and "
    "reviews, Profile (region override), and an Admin catalog screen for admin users."
)
add_p(
    "Video CDN delivery of full movies is explicitly out of scope (same as the "
    "original proposal). Trailers play in-app via YouTube."
)

# ========== 2 ==========
add_h("2. Project Objectives & Scope", 1)
add_h("2.1 Objectives", 2)
add_bullet("Implement core Distributed Database concepts in runnable code (not slides only).")
add_bullet("Match the architecture proposal with a pragmatic, local, free Docker stack.")
add_bullet("Deliver a usable Netflix-like client that exercises sharding, region routing, and recommendations.")
add_bullet("Keep the system demoable on a laptop without paid cloud services.")

add_h("2.2 In Scope", 2)
add_bullet("API gateway + six domain microservices")
add_bullet("Postgres user DB (vertical fragmentation + replication)")
add_bullet("Postgres movie shards US/EU (horizontal sharding + replication + scatter-gather)")
add_bullet("MongoDB watchlist & review shards (polyglot + app-level sharding)")
add_bullet("Redis cache + Redis Streams event bus")
add_bullet("Genre-affinity recommendations; analytics summary endpoint")
add_bullet("React Native Streamix UI (auth, browse, watchlist, reviews, region, admin)")
add_bullet("Manual failover documentation for Postgres movie shards")

add_h("2.3 Out of Scope / Phase-2", 2)
add_bullet("Full Kafka cluster, Milvus/Pinecone vector DB, ClickHouse")
add_bullet("MongoDB mongos sharded cluster (we use application-level sharding)")
add_bullet("Automatic failover orchestration / leader election tooling")
add_bullet("CDN HLS video streaming of full feature films")

# ========== 3 ==========
add_h("3. System Architecture Overview", 1)
add_p(
    "Clients (mobile / future web / admin) talk only to the API Gateway. "
    "The gateway routes by URL path to the correct microservice. Each service "
    "owns its persistence technology."
)
add_p("Logical flow (text diagram):", bold=True)
add_code(
    "StreamixApp\n"
    "    |\n"
    "    v\n"
    "API Gateway :4000\n"
    "    |-- /auth            -> user-service       -> Postgres user (P/R)\n"
    "    |-- /catalog,/admin  -> catalog-service    -> Postgres movies US/EU (P/R) + Redis cache\n"
    "    |-- /watchlist       -> watchlist-service  -> MongoDB watchlist shards\n"
    "    |-- /reviews         -> review-service     -> MongoDB review shards + Redis Streams\n"
    "    |-- /recommendations -> recommendation-svc -> Mongo + Postgres (read)\n"
    "    |-- /analytics       -> analytics-service  -> Redis + Mongo aggregates\n"
)

add_p(
    "This matches the proposal diagram’s layers (Client → Gateway → Services → "
    "specialized stores), implemented with free local containers."
)

# ========== 4 ==========
add_h("4. Technology Stack", 1)
add_table(
    ["Layer", "Technology", "Role"],
    [
        ["Mobile client", "React Native 0.87 (StreamixApp)", "Netflix-style UI"],
        ["API Gateway", "Express + http-proxy-middleware", "Auth/region header forwarding, path routing"],
        ["Services", "Node.js / Express (6 services)", "Domain logic"],
        ["Relational DB", "PostgreSQL 16 (Bitnami legacy images)", "Users + movie catalog"],
        ["Document DB", "MongoDB 7", "Watchlist + reviews"],
        ["Cache / bus", "Redis 7", "TTL cache + Streams events"],
        ["External enrich", "TMDb API", "Movie metadata / trailers"],
        ["Orchestration", "Docker Compose", "Local multi-node cluster"],
        ["Auth", "JWT + bcrypt", "Login/register, admin role"],
    ],
)

# ========== 5 ==========
add_h("5. Distributed Database Concepts Implemented", 1)
add_p(
    "The following concepts are implemented in code and can be demonstrated live."
)
add_table(
    ["Concept", "Implementation evidence"],
    [
        ["Horizontal fragmentation / sharding", "Movies: hash(contentId)%2 → US/EU Postgres. Watchlist: hash(userId). Reviews: hash(movieId) on Mongo."],
        ["Vertical fragmentation", "users_auth vs users_profile in user_schema.sql"],
        ["Primary–replica replication", "User DB + each movie shard; Bitnami streaming replication (WAL)"],
        ["Read/write routing", "Writes → primary; reads → replica (movieShardRouter.js)"],
        ["Scatter-gather queries", "Catalog search/trending fan-out with Promise.allSettled"],
        ["Region-wise distribution", "US/EU labeled shards; X-Client-Region; /catalog/nearby"],
        ["Distribution transparency", "App never picks a host; only gateway + shard routers know topology"],
        ["Failure isolation", "Separate DBs/services; one failed movie shard does not kill search"],
        ["Polyglot persistence", "Postgres + MongoDB + Redis for different workloads"],
        ["Caching", "Redis TTL (~60s) on trending/search"],
        ["Async eventing", "Review publish → Redis Streams → analytics consumer"],
        ["Microservices", "Six independent Node processes + gateway"],
    ],
)

# ========== 6 ==========
add_h("6. Polyglot Persistence Design", 1)
add_p(
    "Polyglot persistence means choosing the right store for each access pattern "
    "instead of forcing everything into one engine."
)
add_h("6.1 Why PostgreSQL for users & movies", 2)
add_bullet("Strong consistency for identity (register/login, roles).")
add_bullet("Mature streaming replication for HA demos.")
add_bullet("Relational catalog with arrays (genres), full-text-friendly fields, aggregates (rating_avg).")

add_h("6.2 Why MongoDB for watchlist & reviews", 2)
add_bullet("Document model fits flexible review payloads and per-user lists.")
add_bullet("Easy horizontal split by key without schema migrations across SQL shards.")
add_bullet("Demonstrates polyglot clearly versus the earlier Postgres-only Phase 1.")

add_h("6.3 Why Redis", 2)
add_bullet("Low-latency catalog cache for hot endpoints.")
add_bullet("Redis Streams as a free local stand-in for Kafka event bus.")
add_bullet("Counters for viva-visible analytics (cache hits/misses, events consumed).")

# ========== 7 ==========
add_h("7. Microservices & API Gateway", 1)
add_h("7.1 Gateway responsibilities", 2)
add_bullet("Listen on :4000 (unchanged for the mobile app).")
add_bullet("Proxy /auth, /catalog, /admin, /watchlist, /reviews, /recommendations, /analytics.")
add_bullet("Forward Authorization and X-Client-Region headers.")
add_bullet("GET /health aggregates each service’s /health (polyglot proof).")

add_h("7.2 Services", 2)
add_table(
    ["Service", "Port (internal)", "Store", "Main routes"],
    [
        ["user-service", "4001", "Postgres user P/R", "/auth/register, /auth/login"],
        ["catalog-service", "4002", "Postgres movies + Redis", "/catalog/*, /admin/*"],
        ["watchlist-service", "4003", "Mongo ×2", "/watchlist"],
        ["review-service", "4004", "Mongo ×2 + Redis Streams", "/reviews"],
        ["recommendation-service", "4005", "Mongo + Postgres reads", "/recommendations"],
        ["analytics-service", "4006", "Redis + Mongo", "/analytics/summary"],
    ],
)

# ========== 8 ==========
add_h("8. Data Model & Fragmentation", 1)
add_h("8.1 Vertical fragmentation (User)", 2)
add_bullet("users_auth — credentials, role, last_login (security-sensitive).")
add_bullet("users_profile — full_name, preferences (presentation data).")
add_bullet("Same logical user, split by column groups; registration writes both in one transaction.")

add_h("8.2 Horizontal fragmentation (Movies)", 2)
add_bullet("Logical table movies split across movie-shard-us-* and movie-shard-eu-*.")
add_bullet("Placement key: content_id → MD5 hash → shard index.")
add_bullet("home_region metadata set at sync time (US/EU label); physical placement is still hash-based.")

add_h("8.3 Horizontal fragmentation (Watchlist / Reviews on Mongo)", 2)
add_bullet("Watchlist items collection sharded by user_id across watchlist-mongo-1/2.")
add_bullet("Reviews collection sharded by movie_id across review-mongo-1/2.")
add_bullet("Responses include shard index and _store: mongodb for demos.")

# ========== 9 ==========
add_h("9. Shard Routing, Replication & Failover", 1)
add_h("9.1 Shard router", 2)
add_p(
    "shared/movieShardRouter.js and shared/mongoShards.js implement the “Shard Router” "
    "box from the architecture diagram: shard map + hash(key) mod N + primary vs replica selection."
)
add_code("shard = md5(key)[0:4] as int  %  N")

add_h("9.2 Replication", 2)
add_bullet("Configured via Bitnami POSTGRESQL_REPLICATION_MODE master/slave.")
add_bullet("Not simulated — genuine Postgres WAL streaming.")
add_bullet("Verify: SELECT * FROM pg_stat_replication; on a movie primary.")

add_h("9.3 Failover demo (manual)", 2)
add_bullet("Documented in backend/failover-demo.md.")
add_bullet("Stop US primary → search still returns EU results (scatter-gather isolation).")
add_bullet("Promote replica with pg_promote(); remap .env addresses — router abstraction preserved.")

# ========== 10 ==========
add_h("10. Region-Aware Routing & Near You", 1)
add_bullet("Mobile detects device timezone (e.g. Asia/Karachi) and maps to nearest of US|EU (Asia → EU).")
add_bullet("Profile screen shows exact location, mapped shard, and active region; user can override.")
add_bullet("Every API call sends X-Client-Region.")
add_bullet("GET /catalog/nearby queries the preferred region’s movie shard first, then falls back.")
add_p(
    "Important viva nuance: there is no Pakistan/Asia physical shard in this cluster — "
    "only US and EU. Exact geo is detected; routing uses the nearest deployed shard."
)

# ========== 11 ==========
add_h("11. Recommendations Engine", 1)
add_p(
    "Content-based genre affinity (not collaborative filtering, not a vector DB)."
)
add_bullet("Load user’s watchlist (Mongo shard) + user’s reviews (scatter across review Mongo shards).")
add_bullet("Look up genres of those movies from Postgres movie shards.")
add_bullet("Score genres (watchlist +1; review weight (rating−3)/2).")
add_bullet("Fetch catalog movies matching top genres; rank by affinity × popularity.")
add_bullet("Cold start (no signals) → trending fallback.")
add_p(
    "Response includes basis string (e.g. “weighted on genres: Action, Drama”) shown on Home."
)

# ========== 12 ==========
add_h("12. Caching & Event Bus (Redis)", 1)
add_h("12.1 Cache", 2)
add_bullet("Keys like catalog:trending:US and hashed search keys; TTL ≈ 60 seconds.")
add_bullet("Logs [cache] HIT / MISS — useful for viva screenshots.")
add_bullet("Counters stats:cache:hits / misses exposed via analytics.")

add_h("12.2 Event bus", 2)
add_bullet("On new review: XADD reviews:events {movieId, userId, rating, shardIndex, ...}.")
add_bullet("Analytics service XREADGROUP consumer acknowledges events.")
add_bullet("Viva line: “Redis Streams plays the Kafka role locally; production would use Kafka.”")

# ========== 13 ==========
add_h("13. Analytics Service", 1)
add_p("GET /analytics/summary returns a viva-friendly snapshot:")
add_bullet("Topology counts (movie / watchlist / review shards)")
add_bullet("Redis cache stats, stream length, events consumed")
add_bullet("Mongo review document counts per shard")
add_bullet("polyglot map: which store holds what")
add_p("This is the ClickHouse stand-in: aggregated stats without a columnar warehouse.")

# ========== 14 ==========
add_h("14. Mobile Application (Streamix)", 1)
add_h("14.1 Screens", 2)
add_table(
    ["Screen", "Features"],
    [
        ["Login / Register", "JWT auth against /auth"],
        ["Home", "Hero banner, Trending, Recommended, Near You, genre rows"],
        ["Search", "Query + genre chips via /catalog/search"],
        ["Watchlist", "List / add / remove via Mongo-backed API"],
        ["Movie Detail", "In-app YouTube trailer, watchlist toggle, star reviews"],
        ["Profile", "Netflix-style account UI, region segment US/EU, admin link, sign out"],
        ["Admin", "TMDb sync / bulk discover (admin role only)"],
    ],
)
add_h("14.2 UX notes", 2)
add_bullet("Netflix red theme (#E50914), dark background.")
add_bullet("Tab icons: Home, Search, Watchlist (bookmark), Profile (person) via Ionicons.")
add_bullet("Region override persists in AsyncStorage.")

# ========== 15 ==========
add_h("15. Admin & Catalog Enrichment (TMDb)", 1)
add_bullet("Admin users have role=admin in users_auth (set via SQL; re-login for JWT).")
add_bullet("POST /admin/movies/sync — pull one title from TMDb, write to hashed movie shard primary.")
add_bullet("POST /admin/movies/bulk-discover — batch by genre/year/language.")
add_bullet("Failure isolation: if TMDb is down, existing local catalog still serves.")

# ========== 16 ==========
add_h("16. Repository Structure", 1)
add_code(
    "distributed-movie-streaming-db/\n"
    "  README.md\n"
    "  backend/\n"
    "    docker-compose.yml\n"
    "    POLYglot-DEMO.md\n"
    "    failover-demo.md\n"
    "    shared/          # auth, redis, pg, mongoShards, movieShardRouter, shardHash\n"
    "    services/        # gateway + 6 microservices\n"
    "    sql/             # Postgres schemas\n"
    "    src/             # tmdbService + compatibility shims\n"
    "    scripts/         # start-all.js, seedMovies.js\n"
    "  StreamixApp/       # Runnable React Native app\n"
    "  mobile-src/        # Source mirror for the RN app\n"
    "  Distributed_Movie_Streaming_System_Architecture.pdf\n"
    "  Distributed Movie Streaming Database Proposal.docx\n"
)

# ========== 17 ==========
add_h("17. How to Run (Setup Guide)", 1)
add_h("17.1 Backend (all-in Docker)", 2)
add_code(
    "cd backend\n"
    "cp .env.example .env   # set TMDB_API_KEY and JWT_SECRET\n"
    "docker compose up -d --build\n"
    "# wait ~30s for replicas\n"
    "curl -s http://localhost:4000/health\n"
)
add_h("17.2 Promote admin", 2)
add_code(
    "docker exec -e PGPASSWORD=app_pass user-db-master psql -U app -d user_db \\\n"
    "  -c \"UPDATE users_auth SET role='admin' WHERE email='you@example.com';\"\n"
    "# Sign out and sign in again in the app\n"
)
add_h("17.3 Mobile", 2)
add_code(
    "cd StreamixApp\n"
    "npm install && cd ios && pod install && cd ..\n"
    "npx react-native run-ios\n"
    "# API_BASE_URL → http://localhost:4000 (simulator)\n"
)

# ========== 18 ==========
add_h("18. Demonstration Scripts (Viva)", 1)
add_p("Use these live demos (also in backend/POLYglot-DEMO.md).")
add_h("18.1 Health / microservices", 2)
add_code("curl -s http://localhost:4000/health | jq")
add_h("18.2 Mongo sharding proof", 2)
add_code(
    "# after login TOKEN=...\n"
    "curl -s -X POST http://localhost:4000/watchlist \\\n"
    "  -H \"Authorization: Bearer $TOKEN\" -H 'Content-Type: application/json' \\\n"
    "  -d '{\"contentId\":27205}' | jq\n"
    "# expect _store: mongodb and shard: 0|1\n"
)
add_h("18.3 Review → Redis Streams", 2)
add_code(
    "curl -s -X POST http://localhost:4000/reviews \\\n"
    "  -H \"Authorization: Bearer $TOKEN\" -H 'Content-Type: application/json' \\\n"
    "  -d '{\"movieId\":27205,\"rating\":5,\"comment\":\"demo\"}' | jq\n"
)
add_h("18.4 Cache HIT/MISS", 2)
add_code(
    "curl -s http://localhost:4000/catalog/trending >/dev/null\n"
    "curl -s http://localhost:4000/catalog/trending >/dev/null\n"
    "docker logs ddb-catalog-service 2>&1 | grep cache | tail\n"
)
add_h("18.5 Analytics / polyglot map", 2)
add_code("curl -s http://localhost:4000/analytics/summary | jq")
add_h("18.6 Postgres replication", 2)
add_code(
    "docker exec -e PGPASSWORD=app_pass movie-shard-us-master \\\n"
    "  psql -U app -d movie_db -c \"SELECT * FROM pg_stat_replication;\"\n"
)

# ========== 19 ==========
add_h("19. Alignment with Architecture Proposal", 1)
add_p(
    "The architecture PDF / proposal is the full target design. This implementation "
    "is a complete Phase-1 backend that covers DDB theory and polyglot microservices "
    "on free local tooling."
)
add_table(
    ["Proposal / diagram item", "Status in this project"],
    [
        ["API Gateway + clients", "Implemented"],
        ["User service + Postgres P/R", "Implemented"],
        ["Catalog sharded + replicated", "Implemented (Postgres US/EU)"],
        ["Watchlist sharded by userId", "Implemented (MongoDB)"],
        ["Reviews sharded by movieId", "Implemented (MongoDB)"],
        ["Recommendation service", "Implemented (genre affinity)"],
        ["Analytics service", "Implemented (Redis+Mongo aggregates)"],
        ["Redis cache", "Implemented"],
        ["Event bus", "Implemented (Redis Streams as Kafka stand-in)"],
        ["Separate microservice processes", "Implemented (6 + gateway)"],
        ["Kafka / Milvus / ClickHouse", "Designed; Phase-2"],
        ["True mongos cluster", "App-level sharding instead"],
        ["CDN full-movie streaming", "Out of scope"],
    ],
)
add_p(
    "Honest viva statement: “The diagram is our complete design. What we run today "
    "implements the DDB core and polyglot microservices locally; Kafka, vector search, "
    "and ClickHouse remain the scale-out Phase-2.”"
)

# ========== 20 ==========
add_h("20. Phase-2 / Future Work", 1)
add_bullet("Replace Redis Streams with Apache Kafka for durable multi-consumer fan-out.")
add_bullet("Add Milvus/Pinecone embeddings for similarity recommendations.")
add_bullet("Add ClickHouse for heavy analytics dashboards.")
add_bullet("Optional MongoDB sharded cluster with mongos for watchlist/reviews.")
add_bullet("Automated failover / health-driven shard-map updates.")
add_bullet("HLS playback if a content CDN is available.")

# ========== 21 ==========
add_h("21. Conclusion & Viva Talking Points", 1)
add_p(
    "Streamix demonstrates a real distributed database system: data is fragmented "
    "(horizontally and vertically), replicated, routed through a shard-aware layer, "
    "queried via scatter-gather, cached, and synchronized asynchronously across "
    "polyglot stores — all behind microservices and a gateway, with a polished mobile UI."
)
add_p("30-second viva script:", bold=True)
add_p(
    "“Our system is a distributed movie catalog. We vertically fragment users and "
    "horizontally shard movies, watchlists, and reviews. Movie shards are region-labeled "
    "US/EU with streaming replicas. A shard router does hash placement and primary/replica "
    "selection. Catalog search is scatter-gather. Near You uses the client region header. "
    "Watchlist and reviews live in MongoDB; Redis caches hot reads and carries review events. "
    "Recommendations are content-based genre affinity. The big architecture diagram also "
    "shows Kafka, Milvus, and ClickHouse — those are Phase-2 scale-out; Phase-1 implements "
    "the distributed database and polyglot microservice backend we can demo on Docker.”"
)

doc.add_page_break()

# ========== APPENDICES ==========
add_h("Appendix A — API Endpoints", 1)
add_table(
    ["Method", "Path", "Service", "Notes"],
    [
        ["POST", "/auth/register", "user", "Creates auth + profile rows"],
        ["POST", "/auth/login", "user", "JWT; may read replica"],
        ["GET", "/catalog/trending", "catalog", "Scatter-gather + Redis cache"],
        ["GET", "/catalog/search", "catalog", "Scatter-gather + Redis cache"],
        ["GET", "/catalog/nearby", "catalog", "Region-preferred shard"],
        ["GET", "/catalog/movie/:id", "catalog", "Single-shard by contentId"],
        ["POST", "/admin/movies/sync", "catalog", "Admin + TMDb"],
        ["POST", "/admin/movies/bulk-discover", "catalog", "Admin + TMDb"],
        ["GET/POST/DELETE", "/watchlist", "watchlist", "Mongo by userId"],
        ["GET", "/reviews/movie/:id", "review", "Mongo by movieId"],
        ["POST", "/reviews", "review", "Mongo + Streams event"],
        ["GET", "/recommendations", "recommendation", "Auth required"],
        ["GET", "/analytics/summary", "analytics", "Viva dashboard JSON"],
        ["GET", "/health", "gateway", "Aggregated service health"],
    ],
)

add_h("Appendix B — Docker Containers", 1)
add_table(
    ["Container", "Role"],
    [
        ["user-db-master / user-db-replica", "Postgres identity P/R"],
        ["movie-shard-us-master / replica", "Movie shard 0 + replica"],
        ["movie-shard-eu-master / replica", "Movie shard 1 + replica"],
        ["watchlist-mongo-1 / 2", "Watchlist app shards"],
        ["review-mongo-1 / 2", "Review app shards"],
        ["ddb-redis", "Cache + Streams"],
        ["ddb-user-service … ddb-analytics-service", "Microservices"],
        ["ddb-api", "API gateway :4000"],
    ],
)

add_h("Appendix C — Key Source Files", 1)
add_table(
    ["Path", "Purpose"],
    [
        ["backend/services/gateway/index.js", "Path proxy gateway"],
        ["backend/shared/movieShardRouter.js", "Postgres movie/user routing"],
        ["backend/shared/mongoShards.js", "Mongo watchlist/review sharding"],
        ["backend/shared/redis.js", "Cache + Streams helpers"],
        ["backend/shared/shardHash.js", "MD5 hash sharding"],
        ["backend/docker-compose.yml", "Full cluster definition"],
        ["backend/POLYglot-DEMO.md", "Demo curl cheat sheet"],
        ["backend/failover-demo.md", "Postgres failover steps"],
        ["StreamixApp/src/screens/*", "Mobile UI"],
        ["StreamixApp/src/services/api.ts", "Client API + X-Client-Region"],
        ["StreamixApp/src/services/region.ts", "Timezone → US/EU mapping"],
    ],
)

add_p("")
end = doc.add_paragraph()
end.alignment = WD_ALIGN_PARAGRAPH.CENTER
er = end.add_run("— End of Report —")
er.italic = True
er.font.size = Pt(11)

doc.save(OUT)
print("Wrote", OUT)
